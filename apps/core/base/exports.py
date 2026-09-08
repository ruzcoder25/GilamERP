"""Admin panelida Excel eksport imkoniyatini beruvchi universal mixin."""

import io

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _get_concrete_fields(model):
    """Modelning DB ustuniga ega barcha fieldlarini qaytaradi."""
    result = []
    for f in model._meta.get_fields():
        if not hasattr(f, "column"):
            continue
        if getattr(f, "many_to_many", False):
            continue
        col_name = getattr(f, "attname", f.name)
        verbose = str(getattr(f, "verbose_name", col_name)).capitalize()
        result.append({"name": col_name, "verbose": verbose})
    return result


def _build_excel(model, ids, selected_fields):
    """Tanlangan fieldlar bo'yicha .xlsx fayl yaratadi va HttpResponse qaytaradi."""
    qs = model.objects.filter(pk__in=ids).values(*selected_fields)

    wb = Workbook()
    ws = wb.active
    ws.title = str(model._meta.verbose_name_plural)[:31]

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    field_map = {f["name"]: f["verbose"] for f in _get_concrete_fields(model)}
    headers = [field_map.get(f, f) for f in selected_fields]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in qs:
        ws.append([("" if row[f] is None else str(row[f])) for f in selected_fields])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{model._meta.model_name}_export.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class ExportExcelMixin:
    """Changelist sahifasida modal orqali Excel eksport qo'shuvchi mixin."""

    change_list_template = "admin/export_excel_changelist.html"

    def get_urls(self):
        from django.urls import path

        meta = self.model._meta
        extra = [
            path(
                "export-excel/",
                self.admin_site.admin_view(self._export_excel_view),
                name=f"{meta.app_label}_{meta.model_name}_export_excel",
            ),
        ]
        return extra + super().get_urls()

    def changelist_view(self, request, extra_context=None):
        """Changelist kontekstiga export maydonlari va URL ni qo'shadi."""
        extra_context = extra_context or {}
        meta = self.model._meta
        extra_context["export_fields"] = _get_concrete_fields(self.model)
        extra_context["export_url"] = reverse(
            f"admin:{meta.app_label}_{meta.model_name}_export_excel"
        )
        return super().changelist_view(request, extra_context=extra_context)

    def _export_excel_view(self, request):
        """POST: ids[] va fields[] qabul qilib Excel fayl qaytaradi."""
        if request.method != "POST":
            return redirect("../")

        ids = request.POST.getlist("ids")
        selected_fields = request.POST.getlist("fields")

        if not ids:
            self.message_user(
                request, "Hech qanday yozuv tanlanmadi.", level=messages.WARNING
            )
            return redirect("../")

        if not selected_fields:
            self.message_user(
                request, "Kamida 1 ta maydon tanlang.", level=messages.WARNING
            )
            return redirect("../")

        return _build_excel(self.model, ids, selected_fields)
